"""Gera os arquivos de relatório (Excel e PDF) a partir de dados simples."""
import io
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

AZUL = "14213D"
FORMATO_DATA = "dd/mm/yyyy hh:mm"


def _txt(dt, formato="%d/%m/%Y %H:%M"):
    return dt.strftime(formato) if dt else "-"


# ------------------------------------------------------------------ Excel
def _folha(ws, titulos, linhas, larguras):
    ws.append(titulos)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for linha in linhas:
        ws.append(linha)
    for linha in ws.iter_rows(min_row=2):
        for c in linha:
            if hasattr(c.value, "year"):
                c.number_format = FORMATO_DATA
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if linhas:
        ws.auto_filter.ref = ws.dimensions


def gerar_excel(p):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "Relatório de operação do estacionamento"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Período"
    ws["B2"] = f"{p['ini']:%d/%m/%Y} a {p['fim']:%d/%m/%Y}"
    ws["A3"] = "Gerado em"
    ws["B3"] = f"{p['gerado_em']:%d/%m/%Y %H:%M} por {p['quem']}"
    ws.append([])
    ws.append(["Indicador", "Valor"])
    for c in ws[5]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
    for rotulo, valor in p["resumo"]:
        ws.append([rotulo, valor])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 34

    _folha(
        wb.create_sheet("Saídas"),
        ["Placa", "Ticket", "Subsolo", "Entrada", "Pago às", "Saída", "Tempo no local", "Minutos no local"],
        [[s["placa"], s["ticket"], s["subsolo"], s["entrada"], s["pago"], s["saida"], s["tempo"], s["minutos"]]
         for s in p["saidas"]],
        [12, 12, 10, 18, 18, 18, 16, 16],
    )
    _folha(
        wb.create_sheet("No estacionamento agora"),
        ["Placa", "Ticket", "Subsolo", "Situação", "Entrada", "Pago às", "Tempo até agora", "Minutos até agora"],
        [[s["placa"], s["ticket"], s["subsolo"], s["situacao"], s["entrada"], s["pago"], s["tempo"], s["minutos"]]
         for s in p["no_momento"]],
        [12, 12, 10, 20, 18, 18, 16, 16],
    )
    _folha(
        wb.create_sheet("Ocupação"),
        ["Subsolo", "Carros", "Capacidade", "Livres", "Situação"],
        [[o["subsolo"], o["carros"], o["capacidade"], o["livres"], o["situacao"]] for o in p["ocupacao"]],
        [14, 10, 12, 10, 14],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
def _tabela(titulos, linhas, larguras, estilos):
    dados = [titulos] + (linhas or [["Nenhum registro."] + [""] * (len(titulos) - 1)])
    t = Table(dados, colWidths=[w * cm for w in larguras], repeatRows=1)
    st = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + AZUL)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c9ced6")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f3f6")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    if not linhas:
        st.append(("SPAN", (0, 1), (-1, 1)))
    t.setStyle(TableStyle(st + estilos))
    return t


def gerar_pdf(p):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.4 * cm, bottomMargin=1.5 * cm,
        title="Relatório de operação do estacionamento",
    )
    estilos = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=estilos["Title"], fontSize=17, alignment=0, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=estilos["Heading2"], fontSize=12, spaceBefore=12, spaceAfter=5)
    pequeno = ParagraphStyle("pq", parent=estilos["Normal"], fontSize=9, textColor=colors.HexColor("#555555"))

    story = [
        Paragraph("Relatório de operação do estacionamento", h1),
        Paragraph(
            f"Período: <b>{p['ini']:%d/%m/%Y}</b> a <b>{p['fim']:%d/%m/%Y}</b> &nbsp;|&nbsp; "
            f"Gerado em {p['gerado_em']:%d/%m/%Y %H:%M} por {escape(p['quem'])}", pequeno),
        Spacer(1, 8),
    ]

    story.append(Paragraph("Resumo", h2))
    story.append(_tabela(["Indicador", "Valor"], [[r, str(v)] for r, v in p["resumo"]], [9, 5],
                         [("ALIGN", (1, 1), (1, -1), "RIGHT")]))

    story.append(Paragraph(f"Carros que já saíram ({len(p['saidas'])})", h2))
    story.append(_tabela(
        ["Placa", "Ticket", "Subsolo", "Entrada", "Pago às", "Saída", "Tempo no local"],
        [[s["placa"], s["ticket"], s["subsolo"], _txt(s["entrada"]), _txt(s["pago"], "%H:%M"),
          _txt(s["saida"]), s["tempo"]] for s in p["saidas"]],
        [3, 3.2, 2.4, 4.2, 2.6, 4.2, 3.6], [],
    ))

    story.append(Paragraph(f"Carros no estacionamento agora ({len(p['no_momento'])})", h2))
    story.append(_tabela(
        ["Placa", "Ticket", "Subsolo", "Situação", "Entrada", "Pago às", "Tempo até agora"],
        [[s["placa"], s["ticket"], s["subsolo"], s["situacao"], _txt(s["entrada"]), _txt(s["pago"], "%H:%M"),
          s["tempo"]] for s in p["no_momento"]],
        [3, 3.2, 2.4, 4.4, 4, 2.6, 3.6], [],
    ))

    story.append(Paragraph("Ocupação por subsolo", h2))
    story.append(_tabela(
        ["Subsolo", "Carros", "Capacidade", "Livres", "Situação"],
        [[o["subsolo"], o["carros"], o["capacidade"], o["livres"], o["situacao"]] for o in p["ocupacao"]],
        [4, 3, 3.5, 3, 3.5], [],
    ))

    def rodape(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawRightString(landscape(A4)[0] - 1.5 * cm, 0.8 * cm, f"Página {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=rodape, onLaterPages=rodape)
    return buf.getvalue()
